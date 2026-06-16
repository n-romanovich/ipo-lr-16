from django.shortcuts import render, get_object_or_404, redirect
from django.contrib.auth.decorators import login_required
from django.contrib.auth import login, update_session_auth_hash
from django.contrib.auth.forms import UserCreationForm, PasswordChangeForm
from django.contrib.auth.views import LoginView
from django.db.models import Q
from django.core.paginator import Paginator
from django.core.mail import EmailMessage
from django.conf import settings
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font
from .models import Product, Category, Manufacturer, Cart, CartItem, Order, OrderItem


def index(request):
    #Последние 6 товаров для блока популярных
    popular_products = Product.objects.all().order_by('-id')[:6]
    categories = Category.objects.all()
    return render(request, 'shop/index.html', {
        'popular_products': popular_products,
        'categories': categories,
    })

#Страница с ссылками на информацию об аторе и лр
def about(request):
    return render(request, 'shop/about.html')

#Страница об авторе лабы
def about_us(request):
    return render(request, 'shop/about_us.html')

#Страница с информацией о ЛР №16-24
def about_shop(request):
    return render(request, 'shop/about_shop.html')

#Каталог с фильтрацией по категории и производителю, а также поиском
def product_list(request):
    items = Product.objects.all()
    categories = Category.objects.all()
    manufacturers = Manufacturer.objects.all()

    #Фильтрация по категории
    category_id = request.GET.get('category')
    if category_id:
        items = items.filter(category_id=category_id)

    #Фильтрация по производителю
    manufacturer_id = request.GET.get('manufacturer')
    if manufacturer_id:
        items = items.filter(manufacturer_id=manufacturer_id)

    #Поиск по названию и описанию
    query = request.GET.get('q')
    if query:
        items = items.filter(Q(name__icontains=query) | Q(description__icontains=query))

    #Пагинация — по 9 товаров на странице
    paginator = Paginator(items, 9)
    page_number = request.GET.get('page')
    items = paginator.get_page(page_number)

    return render(request, 'shop/catalog.html', {
        'items': items,
        'categories': categories,
        'manufacturers': manufacturers,
    })

#Детальная информация о товаре по его ID
def product_detail(request, pk):
    product = get_object_or_404(Product, pk=pk)
    return render(request, 'shop/product_detail.html', {'product': product})

#Добавление товара в корзину (только для авторизованных)
@login_required
def add_to_cart(request, product_id):
    product = get_object_or_404(Product, pk=product_id)
    #Создаем корзину пользователя, если её нет
    cart, created = Cart.objects.get_or_create(user=request.user)

    #Если товар уже в корзине — увеличиваем количество
    cart_item, created = CartItem.objects.get_or_create(cart=cart, product=product, defaults={'quantity': 1})
    if not created:
        if cart_item.quantity < product.stock:
            cart_item.quantity += 1
            cart_item.save()

    return redirect('cart_view')

#Обновление количества товара в корзине
@login_required
def update_cart(request, item_id):
    cart_item = get_object_or_404(CartItem, pk=item_id, cart__user=request.user)

    if request.method == 'POST':
        quantity = int(request.POST.get('quantity', 1))
        #Валидация: количество не должно превышать остаток на складе
        if quantity > cart_item.product.stock:
            quantity = cart_item.product.stock
        if quantity <= 0:
            cart_item.delete()
        else:
            cart_item.quantity = quantity
            cart_item.save()

    return redirect('cart_view')

#Удаление товара из корзины
@login_required
def remove_from_cart(request, item_id):
    cart_item = get_object_or_404(CartItem, pk=item_id, cart__user=request.user)
    cart_item.delete()
    return redirect('cart_view')

#Просмотр корзины пользователя
@login_required
def cart_view(request):
    cart, created = Cart.objects.get_or_create(user=request.user)
    items = cart.items.all()
    total = cart.total_price()
    return render(request, 'shop/cart.html', {'cart': cart, 'items': items, 'total': total})


#Генерация чека в формате Excel
def generate_receipt(order):
    wb = Workbook()
    ws = wb.active
    ws.title = "Чек"

    ws['A1'] = "Чек заказа №" + str(order.id)
    ws['A1'].font = Font(bold=True, size=14)
    ws['A3'] = "Покупатель: " + order.user.username
    ws['A4'] = "Email: " + order.user.email
    ws['A5'] = "Адрес доставки: " + order.address
    ws['A7'] = "Товар"
    ws['B7'] = "Количество"
    ws['C7'] = "Цена"
    ws['D7'] = "Стоимость"

    row = 8
    for item in order.items.all():
        ws.cell(row=row, column=1, value=item.product_name)
        ws.cell(row=row, column=2, value=item.quantity)
        ws.cell(row=row, column=3, value=float(item.price))
        ws.cell(row=row, column=4, value=float(item.item_price()))
        row += 1

    ws.cell(row=row + 1, column=3, value="Итого:")
    ws.cell(row=row + 1, column=4, value=float(order.total_price))
    ws.cell(row=row + 1, column=4).font = Font(bold=True)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


#Отправка чека на email покупателя
def send_receipt_email(order, file_data):
    subject = "Чек заказа №" + str(order.id)
    message = "Спасибо за покупку! Ваш чек во вложении."
    from_email = settings.EMAIL_HOST_USER
    to_email = order.user.email

    email = EmailMessage(subject, message, from_email, [to_email])
    email.attach('cheque_' + str(order.id) + '.xlsx', file_data.read(), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    email.send()


#Оформление заказа
@login_required
def checkout(request):
    cart, created = Cart.objects.get_or_create(user=request.user)
    items = cart.items.all()

    if not items:
        return redirect('product_list')

    if request.method == 'POST':
        address = request.POST.get('address', '')
        if not address:
            return render(request, 'shop/checkout.html', {'items': items, 'error': 'Введите адрес доставки'})

        #Создаем заказ
        order = Order.objects.create(
            user=request.user,
            address=address,
            total_price=cart.total_price()
        )

        #Переносим товары из корзины в заказ
        for cart_item in items:
            OrderItem.objects.create(
                order=order,
                product_name=cart_item.product.name,
                quantity=cart_item.quantity,
                price=cart_item.product.price
            )

        #Генерируем чек и отправляем на email
        receipt = generate_receipt(order)
        send_receipt_email(order, receipt)

        #Очищаем корзину
        cart.items.all().delete()

        return redirect('checkout_success')

    return render(request, 'shop/checkout.html', {'items': items, 'total': cart.total_price()})


#Страница успешного оформления заказа
@login_required
def checkout_success(request):
    return render(request, 'shop/checkout_success.html')


#Вход с Bootstrap-стилями
class PirateLoginView(LoginView):
    def get_form(self, form_class=None):
        form = super().get_form(form_class)
        for f in form.fields.values():
            f.widget.attrs['class'] = 'form-control'
        return form


#Регистрация нового пользователя
def signup(request):
    if request.method == 'POST':
        form = UserCreationForm(request.POST)
        if form.is_valid():
            user = form.save()
            login(request, user)
            return redirect('index')
    else:
        form = UserCreationForm()
    #добавляем Bootstrap классы
    for f in form.fields.values():
        f.widget.attrs['class'] = 'form-control'
    return render(request, 'shop/signup.html', {'form': form})


#Личный кабинет
@login_required
def profile_view(request):
    categories = Category.objects.all()
    return render(request, 'shop/profile.html', {'categories': categories})


#Настройки (смена пароля)
@login_required
def settings_view(request):
    if request.method == 'POST':
        form = PasswordChangeForm(request.user, request.POST)
        if form.is_valid():
            user = form.save()
            update_session_auth_hash(request, user)
            return redirect('profile')
    else:
        form = PasswordChangeForm(request.user)
    for f in form.fields.values():
        f.widget.attrs['class'] = 'form-control'
    return render(request, 'shop/settings.html', {'form': form})

